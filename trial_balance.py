"""
โมดูล logic หลักสำหรับแปลงไฟล์งบทดลอง (.xls) 2 ไฟล์ (เดือนก่อนหน้า + เดือนปัจจุบัน)
ลงในไฟล์ template_ยอดยกมา-ยอดยกไป-.xlsx

หลักการ: โมดูลนี้มีหน้าที่แค่ parse ไฟล์ .xls ต้นทาง (normalize รหัสบัญชี, รวมยอดแถวซ้ำ,
ตรวจับแถวรหัสผิดปกติ) แล้วเขียนดิบๆ ลงชีต "เดือนก่อนหน้า"/"เดือนปัจจุบัน" เท่านั้น
การคำนวณยอดยกไป (ตามหลักเดบิต/เครดิตปกติของหมวดบัญชี) และการจับคู่ยอดยกมาข้ามชีต
ทำโดยสูตร Excel ที่มีอยู่แล้วในเทมเพลต (ดูชีต "วิธีใช้") - ไม่ต้องคำนวณซ้ำใน Python

ไม่ผูกกับ Flask - เรียกใช้/ทดสอบแยกได้ตรงๆ
"""

import os
import re

import xlrd
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(SCRIPT_DIR, "template_ยอดยกมา-ยอดยกไป-.xlsx")

SHEET_PREVIOUS = "เดือนก่อนหน้า"
SHEET_CURRENT = "เดือนปัจจุบัน"

# คอลัมน์ในไฟล์งบทดลองที่นำเข้า (0-indexed, ตาม xlrd)
SRC_CODE_COL = 0
SRC_NAME_COL = 3
SRC_DEBIT_COL = 6          # G: เดบิตเดือนนี้
SRC_CREDIT_COL = 7         # H: เครดิตเดือนนี้
SRC_BALANCE_DEBIT_COL = 9   # J: EndDr (ยอดคงเหลือด้านเดบิต)
SRC_BALANCE_CREDIT_COL = 10  # K: EndCr (ยอดคงเหลือด้านเครดิต)
SRC_DATA_START_ROW = 8
SRC_DATE_ROW = 3
SRC_DATE_COL = 5

# คอลัมน์ในชีต "เดือนก่อนหน้า"/"เดือนปัจจุบัน" ของเทมเพลต (1-indexed, ตาม openpyxl)
TPL_CODE_COL = 1     # A รหัสบัญชี
TPL_NAME_COL = 2     # B ชื่อบัญชี
TPL_DR_COL = 3       # C Dr      <- source G
TPL_CR_COL = 4       # D Cr      <- source H
TPL_ENDDR_COL = 5    # E EndDr   <- source J
TPL_ENDCR_COL = 6    # F EndCr   <- source K
TPL_DATA_START_ROW = 2

ACCOUNT_CODE_PATTERN = re.compile(r"^\d{10}\.\d{3}")

THAI_MONTH_NAMES = [
    "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
]


def thai_month_name(month: int) -> str:
    return THAI_MONTH_NAMES[month - 1]


def parse_thai_date(date_text: str):
    """คืนค่า (be_month, be_year) จากข้อความวันที่ไทยเต็ม เช่น '31 กรกฎาคม 2569'"""
    if not date_text:
        return None
    tokens = [t for t in date_text.strip().split(" ") if t]
    if len(tokens) < 3:
        return None
    month_name = tokens[-2]
    year_text = tokens[-1]
    try:
        year = int(year_text)
    except ValueError:
        return None
    if month_name not in THAI_MONTH_NAMES:
        return None
    be_month = THAI_MONTH_NAMES.index(month_name) + 1
    return be_month, year


def _cell_str(sheet, row, col):
    try:
        v = sheet.cell_value(row, col)
    except IndexError:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() if v is not None else ""


def _cell_num(sheet, row, col):
    try:
        v = sheet.cell_value(row, col)
    except IndexError:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def parse_source_file(file_path_or_stream):
    """
    อ่านไฟล์งบทดลอง .xls แล้ว normalize รหัสบัญชี + รวมยอด (แยก EndDr/EndCr ไว้คนละคอลัมน์
    ไม่ผสมกันเป็นค่าเดียว เพราะสูตรในเทมเพลตจะเป็นคนคำนวณยอดยกไปที่ถูกต้องตามหมวดบัญชีเอง)

    คืนค่า dict:
        be_month, be_year: รอบบัญชีที่อ่านได้จากเนื้อไฟล์
        parsed: {รหัสบัญชี(14 ตัวอักษร): {"name":.., "debit":.., "credit":.., "end_dr":.., "end_cr":..}}
        malformed: [{"row":.., "code":.., "name":.., "debit":.., "credit":.., "end_dr":.., "end_cr":..}, ...]
            แถวที่รหัสบัญชีอ่านไม่ออกแต่มีมูลค่าเงินติดมาด้วย (ห้ามทิ้งเงียบๆ)
    """
    # ไฟล์เหล่านี้มักไม่มี CODEPAGE record ต้องระบุ cp874 (Windows Thai) เองมิฉะนั้น
    # ข้อความไทยจะถูกอ่านผิดเพี้ยน (xlrd จะ fallback ไปเป็น iso-8859-1)
    if hasattr(file_path_or_stream, "read"):
        wb = xlrd.open_workbook(file_contents=file_path_or_stream.read(), encoding_override="cp874")
    else:
        wb = xlrd.open_workbook(file_path_or_stream, encoding_override="cp874")
    sheet = wb.sheet_by_index(0)

    date_text = _cell_str(sheet, SRC_DATE_ROW, SRC_DATE_COL)
    parsed_date = parse_thai_date(date_text)
    if parsed_date is None:
        raise ValueError(
            "ไม่พบวันที่ของรอบบัญชีในไฟล์ (แถวที่ {0}, ข้อความที่อ่านได้: '{1}') "
            "กรุณาตรวจสอบรูปแบบไฟล์".format(SRC_DATE_ROW + 1, date_text)
        )
    be_month, be_year = parsed_date

    parsed = {}
    malformed = []

    for i in range(SRC_DATA_START_ROW, sheet.nrows):
        raw = _cell_str(sheet, i, SRC_CODE_COL)
        if not raw:
            continue

        debit = _cell_num(sheet, i, SRC_DEBIT_COL)
        credit = _cell_num(sheet, i, SRC_CREDIT_COL)
        # ไฟล์ต้นฉบับเก็บ EndCr เป็นค่าติดลบเมื่อบัญชีมียอดคงเหลือด้านเครดิต (ยืนยันจากข้อมูลจริง
        # เช่น "เจ้าหนี้การค้า" ได้ EndCr=-26000) แต่สูตรในเทมเพลตต้องการ EndDr/EndCr เป็นค่าบวก
        # (นำหน้าไปลบกันเองตามหมวดบัญชีใน G) จึงต้องแปลงเป็นค่าสัมบูรณ์ก่อนเก็บ
        end_dr = abs(round(_cell_num(sheet, i, SRC_BALANCE_DEBIT_COL), 2))
        end_cr = abs(round(_cell_num(sheet, i, SRC_BALANCE_CREDIT_COL), 2))

        m = ACCOUNT_CODE_PATTERN.match(raw)
        if not m:
            if debit != 0 or credit != 0 or end_dr != 0 or end_cr != 0:
                malformed.append({
                    "row": i + 1,
                    "code": raw,
                    "name": _cell_str(sheet, i, SRC_NAME_COL),
                    "debit": debit,
                    "credit": credit,
                    "end_dr": end_dr,
                    "end_cr": end_cr,
                })
            continue

        code = raw[:14]
        if code not in parsed:
            parsed[code] = {"debit": 0.0, "credit": 0.0, "end_dr": 0.0, "end_cr": 0.0,
                             "name": _cell_str(sheet, i, SRC_NAME_COL)}
        parsed[code]["debit"] += debit
        parsed[code]["credit"] += credit
        parsed[code]["end_dr"] += end_dr
        parsed[code]["end_cr"] += end_cr

    return {
        "be_month": be_month,
        "be_year": be_year,
        "parsed": parsed,
        "malformed": malformed,
    }


def diff_codes(previous_parsed, current_parsed):
    """
    หารหัสบัญชีที่มีอยู่ในไฟล์หนึ่งแต่ไม่มีในอีกไฟล์ (เสริมจาก column "ตรวจสอบ" ในเทมเพลต
    ให้เห็นในหน้าเว็บก่อนดาวน์โหลดเลย ไม่ต้องเปิด Excel ก่อน)

    คืนค่า: (only_in_previous, only_in_current) - แต่ละอันเป็น list ของ "รหัส - ชื่อ"
    """
    prev_codes = set(previous_parsed.keys())
    cur_codes = set(current_parsed.keys())

    only_in_previous = []
    for code in sorted(prev_codes - cur_codes):
        name = previous_parsed[code].get("name", "")
        only_in_previous.append(code + (" - " + name if name else ""))

    only_in_current = []
    for code in sorted(cur_codes - prev_codes):
        name = current_parsed[code].get("name", "")
        only_in_current.append(code + (" - " + name if name else ""))

    return only_in_previous, only_in_current


def _write_month_sheet(ws, parsed):
    """เขียนข้อมูลที่ parse แล้วลงชีต (เดือนก่อนหน้า/เดือนปัจจุบัน) เริ่มแถว 2 เรียงตามรหัสบัญชี
    เขียนแค่ A-F เท่านั้น ไม่แตะคอลัมน์ G/I ที่เป็นสูตรคำนวณอัตโนมัติของเทมเพลต"""
    row = TPL_DATA_START_ROW
    for code in sorted(parsed.keys()):
        v = parsed[code]
        ws.cell(row=row, column=TPL_CODE_COL).value = code
        ws.cell(row=row, column=TPL_NAME_COL).value = v.get("name", "")
        ws.cell(row=row, column=TPL_DR_COL).value = round(v["debit"], 2)
        ws.cell(row=row, column=TPL_CR_COL).value = round(v["credit"], 2)
        ws.cell(row=row, column=TPL_ENDDR_COL).value = round(v["end_dr"], 2)
        ws.cell(row=row, column=TPL_ENDCR_COL).value = round(v["end_cr"], 2)
        row += 1


def build_output_workbook(previous_parsed, current_parsed, template_path: str = TEMPLATE_PATH):
    """
    สร้าง Workbook (openpyxl) จาก template โดยเติมข้อมูลดิบลงชีต "เดือนก่อนหน้า" และ "เดือนปัจจุบัน"
    ชีต "ผลลัพธ์" และ "ปิดบัญชีปลายปีงบ(ก.ย.->ต.ค.)" เป็นสูตรอยู่แล้วในเทมเพลต ไม่ต้องแตะ
    """
    wb = openpyxl.load_workbook(template_path)
    _write_month_sheet(wb[SHEET_PREVIOUS], previous_parsed)
    _write_month_sheet(wb[SHEET_CURRENT], current_parsed)
    return wb


def output_filename(be_month, be_year):
    return "งบทดลอง {0} {1}.xlsx".format(thai_month_name(be_month), be_year)
